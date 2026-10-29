"""Generate fictional supplier documents and a gold-standard dataset for BidLens."""
from pathlib import Path
import csv, io, random
import pandas as pd
from openpyxl import Workbook
from docx import Document
from reportlab.lib.pagesizes import letter, landscape
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent
DOCS = ROOT / "data" / "source_documents"
DOCS.mkdir(parents=True, exist_ok=True)

items=[]
for family, prefix, size in [("Small box","S","Small"),("Medium box","M","Medium"),("Large box","L","Large")]:
    for n in range(1,11):
        ix=len(items)+1
        items.append({"item_id":f"BX-{ix:02d}","item":f"{family} {n:02d}","family":family,"description":f"{size} shipping box","quantity":[500,750,1000][(n-1)%3],"required_delivery_days":7})

suppliers = [
    {"id":"SUP-A","name":"Deccan Packaging (Sample)","format":"Excel workbook","file":"supplier_a_deccan_packaging.xlsx"},
    {"id":"SUP-B","name":"Charminar Boxes (Sample)","format":"PDF quote","file":"supplier_b_charminar_boxes.pdf"},
    {"id":"SUP-C","name":"Musi Cartons (Sample)","format":"Scanned image","file":"supplier_c_musi_cartons_scan.png"},
    {"id":"SUP-D","name":"Metro Pack Hyderabad (Sample)","format":"Email response","file":"supplier_d_metro_pack.eml"},
    {"id":"SUP-E","name":"Telangana Box Works (Sample)","format":"Word document","file":"supplier_e_telanganaboxworks.docx"},
]

all_quotes=[]
for s_idx, s in enumerate(suppliers):
    for ix, item in enumerate(items):
        missing = ((s_idx==1 and ix in {7,18,27}) or (s_idx==3 and ix%3!=0) or (s_idx==4 and ix in {4,14,24}))
        ambiguous = (s_idx==2 and ix==12)
        base = 8.35 + (len(item["description"])*0.13) + (ix%5)*0.42
        price = round(base * [1.0,0.96,0.93,0.91,0.98][s_idx], 2)
        currency = "USD" if s_idx==4 else "INR"
        basis = "per 100 items" if s_idx==4 and ix%2==0 else "per item"
        if s_idx==4 and basis=="per 100 items": price=round(price*100/83,2) # genuine USD per 100, normalized later at 83 INR/USD
        if s_idx==4 and basis=="per item": price=round(price/83,2)
        delivery_days = [5,6,4,7,8][s_idx]
        delivery_charge = [0,250,300,None,350][s_idx]
        if s_idx==3: delivery_charge=None
        if missing or ambiguous: price_value=None
        else: price_value=price
        issue=[]
        if missing: issue.append("No quote for this requested item")
        if ambiguous: issue.append("Scanned price is ambiguous; manual confirmation required")
        if s_idx==1 and ix==0: issue.append("PDF notes mention a discount; applicability is unconfirmed")
        if s_idx==3: issue.append("Freight/additional charges not specified")
        if s_idx==4: issue.append("USD quote; conversion assumption required")
        all_quotes.append({**item,"supplier_id":s["id"],"supplier":s["name"],"source_format":s["format"],"source_file":s["file"],"quoted_price":price_value,"original_price_text":"8.?" if ambiguous else ("" if missing else f"{price:.2f}"),"currency":currency,"price_basis":basis,"delivery_days":None if missing else delivery_days,"delivery_charge":delivery_charge,"payment_terms":"Net 30 days" if s_idx in (0,2,4) else ("Net 45 days" if s_idx==1 else "Not specified"),"quote_validity_days":30 if s_idx!=3 else None,"discount_note":"10% volume discount mentioned in notes; minimum quantity and applicability not stated" if s_idx==1 else "","issue":"; ".join(issue),"source_evidence": ""})

quotes_df=pd.DataFrame(all_quotes)
items_df=pd.DataFrame(items)
items_df.to_csv(ROOT/"data"/"rfx_items.csv",index=False)
quotes_df.to_csv(ROOT/"data"/"gold_standard_quotes.csv",index=False)

# Helper: only suppliers' actually stated rows are included in documents.
def rows_for(s_idx):
    d=quotes_df[quotes_df.supplier_id==suppliers[s_idx]["id"]].copy()
    if s_idx==1: d=d[d.item_id.isin([x["item_id"] for i,x in enumerate(items) if i not in {7,18,27}])]
    if s_idx==3: d=d[d.item_id.isin([x["item_id"] for i,x in enumerate(items) if i%3==0])]
    if s_idx==4: d=d[d.item_id.isin([x["item_id"] for i,x in enumerate(items) if i not in {4,14,24}])]
    return d

def row_text(r, include_supplier=False):
    price = r["original_price_text"] if r["original_price_text"] else "NOT QUOTED"
    # The scan explicitly carries an ambiguous value, rather than hiding it as missing.
    return f"{r['item_id']} | {r['item']} | Qty {r['quantity']} | Price {price} | {r['currency']} | {r['price_basis']} | Delivery {r['delivery_days'] if pd.notna(r['delivery_days']) else 'NOT STATED'} days | Freight {r['delivery_charge'] if pd.notna(r['delivery_charge']) else 'NOT STATED'} | Payment {r['payment_terms']}"

# Supplier A: complete Excel quote.
d=rows_for(0)
with pd.ExcelWriter(DOCS/suppliers[0]["file"], engine="openpyxl") as w:
    out=d[["item_id","item","quantity","original_price_text","currency","price_basis","delivery_days","delivery_charge","payment_terms","quote_validity_days"]].copy()
    out.columns=["Item ID","Item description","Quantity offered","Unit price","Currency","Price basis","Delivery days","Freight / delivery charge","Payment terms","Quote validity (days)"]
    out.to_excel(w,index=False,sheet_name="Quotation")
    pd.DataFrame([{"Supplier":"Deccan Packaging (Sample)","Quote date":"2026-10-28","Delivery location":"Hyderabad, Telangana","Notes":"All prices in INR. Freight shown per quotation. GST is included in quoted prices."}]).to_excel(w,index=False,sheet_name="Cover")

# Supplier B: PDF with missing rows and an unqualified discount note.
styles=getSampleStyleSheet()
styles.add(ParagraphStyle(name="SmallText", parent=styles["BodyText"], fontSize=7, leading=9))
pdf_path=DOCS/suppliers[1]["file"]
d=rows_for(1)
story=[Paragraph("CHARMINAR BOXES (SAMPLE) — QUOTATION",styles["Title"]),Paragraph("Quote date: 28 October 2026 | Delivery location: Hyderabad, Telangana",styles["BodyText"]),Spacer(1,8),Paragraph("Commercial note: 10% volume discount mentioned. Minimum order quantity and applicability across line items are not specified; please confirm before applying.",styles["BodyText"]),Paragraph("Tax treatment: not stated in this quotation.",styles["BodyText"]),Spacer(1,8)]
data=[["Item ID","Description","Qty","Price","Currency","Basis","Delivery","Freight"]]
for _,r in d.iterrows():
    data.append([r.item_id,r["item"],str(r.quantity),r.original_price_text,r.currency,r.price_basis,str(int(r.delivery_days)) if pd.notna(r.delivery_days) else "Not stated",str(int(r.delivery_charge)) if pd.notna(r.delivery_charge) else "Not stated"])
t=Table(data,repeatRows=1,colWidths=[48,110,38,48,42,65,45,48])
t.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#e9eef5")),("GRID",(0,0),(-1,-1),0.35,colors.HexColor("#aab2bd")),("FONTSIZE",(0,0),(-1,-1),6.5), ("VALIGN",(0,0),(-1,-1),"TOP"), ("LEFTPADDING",(0,0),(-1,-1),3),("RIGHTPADDING",(0,0),(-1,-1),3)]))
story.append(t)
story.append(PageBreak())
story.append(Paragraph("Structured text copy for accessible quotation review", styles["Heading2"]))
story.append(Paragraph("The lines below preserve the quoted fields in a consistent text representation for extraction and audit testing.", styles["BodyText"]))
for _, rr in d.iterrows():
    story.append(Paragraph(row_text(rr), styles["SmallText"]))
SimpleDocTemplate(str(pdf_path),pagesize=landscape(letter),rightMargin=20,leftMargin=20,topMargin=24,bottomMargin=24).build(story)

# Supplier C: a long scan-like image. Text transcript is included as a reproducible demo OCR fallback.
d=rows_for(2)
scan_lines=["MUSI CARTONS (SAMPLE) - QUOTATION", "Quote date: 28 October 2026 | Deliver to Hyderabad", "All quoted prices in INR per item; delivery charge per quotation: INR 300", "Tax treatment: not stated in scan.", "Item ID | Description | Qty | Price | Currency | Basis | Delivery days"]
for _,r in d.iterrows():
    p="8.?" if r.item_id=="BX-13" else r.original_price_text
    scan_lines.append(f"{r.item_id} | {r["item"]} | {r.quantity} | {p} | INR | per item | {int(r.delivery_days)}")
scan_lines += ["Payment terms: Net 30 days", "Note: Price for BX-13 is handwritten and unclear; supplier confirmation required."]
transcript="\n".join(scan_lines)
(DOCS/"supplier_c_musi_cartons_scan_ocr_reference.txt").write_text(transcript,encoding="utf-8")
try: font=ImageFont.truetype("DejaVuSansMono.ttf",20)
except: font=ImageFont.load_default()
width=1800; line_h=34; height=90+line_h*len(scan_lines)
im=Image.new("RGB",(width,height),(250,248,239)); draw=ImageDraw.Draw(im)
draw.text((48,24),"SUPPLIER QUOTATION — SCANNED COPY",font=font,fill=(30,30,30))
for i,line in enumerate(scan_lines):
    draw.text((48,68+i*line_h),line,font=font,fill=(35,35,35))
im.save(DOCS/suppliers[2]["file"])

# Supplier D: email with only every third item priced and charges not stated.
d=rows_for(3)
email_lines=["From: sales@metro-pack.example","To: procurement@example.invalid","Date: 28 October 2026","Subject: Re: Shipping box quotation","","Hello Procurement Team,","","We can quote the following requested items. Prices are in INR per item. Delivery is approximately 7 days after order confirmation. Freight, taxes, and other additional charges are not specified in this email. Please confirm the final charge treatment with us.",""]
for _,r in d.iterrows(): email_lines.append(row_text(r))
email_lines += ["","Payment terms: Not specified.","Regards,","Metro Pack Hyderabad (Sample)"]
(DOCS/suppliers[3]["file"]).write_text("\n".join(email_lines),encoding="utf-8")

# Supplier E: Word document with USD and mixed per-item / per-100 price bases.
d=rows_for(4)
doc=Document(); doc.add_heading("Telangana Box Works (Sample) — Quotation",0)
doc.add_paragraph("Quote date: 28 October 2026 | Delivery location: Hyderabad, Telangana")
doc.add_paragraph("Prices are quoted in USD. Price basis varies by line and is shown explicitly. Delivery charge: USD 350 per quotation. Delivery lead time: 8 days. Payment terms: Net 30 days.")
doc.add_paragraph("Currency conversion is not supplied by the vendor. Any INR conversion must use a separately disclosed buyer-approved assumption. Tax treatment is not stated in this quotation.")
table=doc.add_table(rows=1, cols=7); table.style="Table Grid"
for cell,val in zip(table.rows[0].cells,["Item ID","Description","Qty","Quoted price","Currency","Price basis","Delivery days"]): cell.text=val
for _,r in d.iterrows():
    vals=[r.item_id,r["item"],str(r.quantity),r.original_price_text,r.currency,r.price_basis,str(int(r.delivery_days))]
    cells=table.add_row().cells
    for c,val in zip(cells,vals): c.text=str(val)
doc.add_paragraph("Quote validity: 30 days. All amounts are synthetic examples for demonstration only.")
doc.save(DOCS/suppliers[4]["file"])

# Record the supplier-level tax treatment explicitly in the gold-standard dataset.
for idx, row in quotes_df.iterrows():
    if row["supplier_id"] == "SUP-A": quotes_df.loc[idx,"tax_status"] = "GST included in quoted prices"
    else: quotes_df.loc[idx,"tax_status"] = "Not stated in source document"

# Build source evidence pointers into gold standard.
for s in suppliers:
    s_df=quotes_df[quotes_df.supplier_id==s["id"]]
    for idx,r in s_df.iterrows():
        if s["id"]=="SUP-A": ev=f"Excel sheet 'Quotation', row {int(r.name)+2}: {r.item_id}"
        elif s["id"]=="SUP-B": ev=f"PDF table row for {r.item_id}; page varies by layout"
        elif s["id"]=="SUP-C": ev=f"Scanned quotation image, line for {r.item_id}; OCR reference transcript packaged"
        elif s["id"]=="SUP-D": ev=f"Email body line for {r.item_id}"
        else: ev=f"Word quotation table row for {r.item_id}"
        quotes_df.loc[idx,"source_evidence"]=ev
quotes_df.to_csv(ROOT/"data"/"gold_standard_quotes.csv",index=False)
with (ROOT/"data"/"supplier_manifest.json").open("w",encoding="utf-8") as f:
    import json; json.dump(suppliers,f,indent=2)
print(f"Generated {len(items)} items, {len(quotes_df)} gold-standard quote rows, and {len(suppliers)} source documents in {DOCS}")
